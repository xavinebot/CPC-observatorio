"""Pellet en Austria (proPellets, índice PPI06 en XLSX) y Alemania (C.A.R.M.E.N. e.V., tabla HTML mensual).

- proPellets: el XLSX está en una ruta que cambia cada año (`/assets/upload/Pelletpreise/<año>/PPI06.xlsx`); se
  localiza el enlace en la página. Hoja única: fila "Jahr" con los años en columnas y una fila por mes.
- C.A.R.M.E.N.: la tabla histórica (wpDataTables) viene incrustada en el HTML con columnas
  `Zahl_Monat, Jahr, Jahr_Monat, Sackware, 2 Tonnen, 5 Tonnen, ...` en €/t con IVA. Usamos "5 Tonnen".
"""
from __future__ import annotations

import io
import re

import openpyxl
from bs4 import BeautifulSoup

from ..util import http_get, save_raw
from .base import Collector as _Base

PP_PAGE = "https://www.propellets.at/aktuelle-pelletpreise"
CARMEN_PAGE = "https://www.carmen-ev.de/service/marktueberblick/marktpreise-energieholz/marktpreise-pellets/"
UA = {"User-Agent": "Mozilla/5.0 (compatible; CPC-Observatorio/0.1; +https://cristalesparachimeneas.es/contacto/)"}
MONTHS_DE = ["jän", "feb", "mär", "apr", "mai", "jun", "jul", "aug", "sep", "okt", "nov", "dez"]


class Collector(_Base):
    name = "pellets_eu"
    min_records = 200

    def fetch(self, *, backfill: bool = False) -> dict[str, list[tuple]]:
        out = {}
        try:
            out["pellet_indice_at"] = fetch_propellets()
        except Exception as e:  # noqa: BLE001
            print(f"   pellets_eu: proPellets no leído ({e})")
        try:
            out["pellet_granel_de"] = fetch_carmen()
        except Exception as e:  # noqa: BLE001
            print(f"   pellets_eu: C.A.R.M.E.N. no leído ({e})")
        if not out:
            raise RuntimeError("pellets_eu: ninguna de las dos fuentes respondió")
        return out


def fetch_propellets() -> list[tuple]:
    html = http_get(PP_PAGE, headers=UA).text
    m = re.search(r'href="(https?://[^"#]*PPI06\.xlsx)', html)
    if not m:
        raise RuntimeError("proPellets: no encuentro el enlace a PPI06.xlsx")
    content = http_get(m.group(1), headers=UA).content
    save_raw("propellets_PPI06.xlsx", content)
    return parse_propellets(content)


def parse_propellets(content: bytes) -> list[tuple]:
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    years = None
    pts = []
    for r in ws.iter_rows(values_only=True):
        cells = list(r)
        label = next((str(c).strip().lower() for c in cells if isinstance(c, str)), "")
        if label == "jahr":
            years = {i: int(c) for i, c in enumerate(cells) if isinstance(c, (int, float))}
            continue
        if years and label[:3] in MONTHS_DE:
            month = MONTHS_DE.index(label[:3]) + 1
            for i, y in years.items():
                v = cells[i] if i < len(cells) else None
                if isinstance(v, (int, float)):
                    pts.append((f"{y}-{month:02d}-01", float(v)))
    if not pts:
        raise RuntimeError("proPellets: hoja sin datos reconocibles")
    return sorted(pts)


def fetch_carmen() -> list[tuple]:
    html = http_get(CARMEN_PAGE, headers=UA).text
    save_raw("carmen_pellets.html", html)
    return parse_carmen(html)


def parse_carmen(html: str) -> list[tuple]:
    soup = BeautifulSoup(html, "html.parser")
    for t in soup.find_all("table"):
        hdr = [th.get_text(strip=True) for th in t.find_all("th")]
        if "Jahr_Monat" in hdr and "5 Tonnen" in hdr:
            i_m, i_y, i_v = hdr.index("Zahl_Monat"), hdr.index("Jahr"), hdr.index("5 Tonnen")
            pts = []
            for tr in t.find_all("tr"):
                td = [x.get_text(strip=True) for x in tr.find_all("td")]
                if len(td) <= i_v or not td[i_y].isdigit():
                    continue
                v = td[i_v].replace(".", "").replace(",", ".")
                if not v:
                    continue
                pts.append((f"{int(td[i_y])}-{int(td[i_m]):02d}-01", float(v) / 1000.0))  # €/t → €/kg
            if pts:
                return sorted(pts)
    raise RuntimeError("C.A.R.M.E.N.: no encuentro la tabla histórica de pellets")
